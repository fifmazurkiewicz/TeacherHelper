from __future__ import annotations

import io

from docx import Document

from teacher_helper.infrastructure.export import text_to_docx


def test_text_to_docx_renders_markdown_as_styled_word_content() -> None:
    raw = """# Materiał

## Najważniejsze pojęcia

- Pierwszy **ważny** punkt
- Drugi punkt z `terminem`

1. Zadanie pierwsze

> Wskazówka dla nauczyciela
"""

    result = text_to_docx(raw, title="Materiał")
    doc = Document(io.BytesIO(result))

    texts = [paragraph.text for paragraph in doc.paragraphs]
    assert texts.count("Materiał") == 1
    assert "Najważniejsze pojęcia" in texts
    assert "Pierwszy ważny punkt" in texts
    assert "Drugi punkt z terminem" in texts
    assert "Wskazówka dla nauczyciela" in texts
    assert not any("**" in text or "`" in text for text in texts)

    heading = next(p for p in doc.paragraphs if p.text == "Najważniejsze pojęcia")
    assert heading.style.name == "Heading 2"
    bullet = next(p for p in doc.paragraphs if p.text == "Pierwszy ważny punkt")
    assert bullet.style.name == "List Bullet"
    assert any(run.bold and run.text == "ważny" for run in bullet.runs)
    assert doc.sections[0].footer.paragraphs[0].text == "TeacherHelper"


def test_text_to_docx_keeps_plain_paragraphs_readable() -> None:
    result = text_to_docx("Pierwszy akapit.\n\nDrugi akapit.", title="Koncepcja lekcji")
    doc = Document(io.BytesIO(result))

    assert [p.text for p in doc.paragraphs if p.text in {"Pierwszy akapit.", "Drugi akapit."}] == [
        "Pierwszy akapit.",
        "Drugi akapit.",
    ]
    assert doc.styles["Normal"].font.size.pt == 10.5
