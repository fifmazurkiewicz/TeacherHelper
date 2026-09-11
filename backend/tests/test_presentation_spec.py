from __future__ import annotations

import base64
import io
import zipfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from teacher_helper.infrastructure.presentation_spec import (
    extract_pptx_slide_images,
    normalize_presentation_spec,
    pptx_to_spec,
    pptx_to_spec_and_images,
    spec_to_pptx_bytes,
)


_PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def test_normalize_preserves_layout_and_speaker_notes() -> None:
    spec = normalize_presentation_spec(
        {
            "title": "Lekcja",
            "description": "Cel",
            "slides": [
                {
                    "title": "Porównanie",
                    "bullets": ["A", "B"],
                    "layout": "comparison",
                    "speaker_notes": "Zapytaj klasę o różnice.",
                },
                {"title": "Błędny układ", "bullets": [], "layout": "cards"},
            ],
        }
    )

    assert spec is not None
    assert spec["slides"][0]["layout"] == "comparison"
    assert spec["slides"][0]["speaker_notes"] == "Zapytaj klasę o różnice."
    assert spec["slides"][1]["layout"] == "text"


def test_pptx_roundtrip_keeps_embedded_image_and_notes() -> None:
    spec = normalize_presentation_spec(
        {
            "title": "Obieg wody",
            "description": "Klasa 4",
            "slides": [
                {
                    "title": "Parowanie",
                    "bullets": ["Słońce ogrzewa wodę"],
                    "include_image": True,
                    "layout": "image_right",
                    "speaker_notes": "Poproś o przykład z życia.",
                }
            ],
        }
    )
    assert spec is not None

    deck = spec_to_pptx_bytes(spec, slide_images={0: _PNG_1X1})
    images = extract_pptx_slide_images(deck)
    parsed = pptx_to_spec(deck)

    assert images[0] == _PNG_1X1
    assert parsed is not None
    assert parsed["slides"][0]["include_image"] is True
    assert "Poproś o przykład" in parsed["slides"][0]["speaker_notes"]
    assert parsed["theme"]["background"] == "#1A243A"


def test_embedded_image_stays_inside_content_box() -> None:
    spec = normalize_presentation_spec(
        {
            "title": "Test",
            "description": "",
            "slides": [
                {
                    "title": "Ilustracja",
                    "bullets": ["Punkt"],
                    "include_image": True,
                    "layout": "image_right",
                }
            ],
        }
    )
    assert spec is not None
    prs = Presentation(io.BytesIO(spec_to_pptx_bytes(spec, slide_images={0: _PNG_1X1})))
    picture = next(shape for shape in prs.slides[1].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)

    assert picture.left >= Inches(6.25)
    assert picture.top >= Inches(1.25)
    assert picture.left + picture.width <= Inches(12.75)
    assert picture.top + picture.height <= Inches(6.6)


def test_comparison_layout_uses_two_content_columns() -> None:
    spec = normalize_presentation_spec(
        {
            "title": "Test",
            "description": "",
            "slides": [
                {
                    "title": "Ssaki i ptaki",
                    "bullets": ["Ssaki karmią mlekiem", "Ssaki mają sierść", "Ptaki składają jaja", "Ptaki mają pióra"],
                    "layout": "comparison",
                }
            ],
        }
    )
    assert spec is not None
    prs = Presentation(io.BytesIO(spec_to_pptx_bytes(spec)))
    content = [shape.text for shape in prs.slides[1].placeholders if shape.has_text_frame and shape != prs.slides[1].shapes.title]

    assert any("Ssaki" in text for text in content)
    assert any("Ptaki" in text for text in content)

    parsed = pptx_to_spec(prs_to_bytes(prs))
    assert parsed is not None
    assert parsed["slides"][0]["layout"] == "comparison"
    assert parsed["slides"][0]["bullets"] == [
        "Ssaki karmią mlekiem",
        "Ssaki mają sierść",
        "Ptaki składają jaja",
        "Ptaki mają pióra",
    ]


def prs_to_bytes(prs: Presentation) -> bytes:
    stream = io.BytesIO()
    prs.save(stream)
    return stream.getvalue()


def test_comparison_layout_disables_overlapping_image() -> None:
    spec = normalize_presentation_spec(
        {
            "title": "Test",
            "description": "",
            "slides": [
                {
                    "title": "Porównanie",
                    "bullets": ["A", "B"],
                    "layout": "comparison",
                    "include_image": True,
                    "image_hint": "Diagram",
                }
            ],
        }
    )
    assert spec is not None
    assert spec["slides"][0]["include_image"] is False
    assert spec["slides"][0]["image_hint"] is None


def test_rejects_pptx_with_unsafe_compression_ratio() -> None:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"0" * 1_000_000)

    spec, images = pptx_to_spec_and_images(stream.getvalue())

    assert spec is None
    assert images == {}
